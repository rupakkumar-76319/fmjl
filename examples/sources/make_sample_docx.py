import sys
from docx import Document
from docx.shared import Cm
from docx.enum.text import WD_BREAK

out, image = sys.argv[1], sys.argv[2]
d = Document()
d.core_properties.title = "Solar Schools Programme: Maintenance Guide"
d.core_properties.author = "Rupak Kumar"
d.sections[0].header.paragraphs[0].text = "Solar Schools Programme - Maintenance Guide"
d.sections[0].footer.paragraphs[0].text = "Internal document"

d.add_heading("Solar Schools Programme: Maintenance Guide", level=0)
p = d.add_paragraph("This guide tells school staff how to keep a rooftop solar installation working. ")
p.add_run("Read it before the first cleaning.").bold = True
d.add_heading("1. Monthly tasks", level=1)
d.add_paragraph("Do these tasks once a month, on a dry day:")
d.add_paragraph("Clean every panel with water and a soft cloth", style="List Bullet")
d.add_paragraph("Check that no cable is loose or exposed", style="List Bullet")
d.add_paragraph("Write the inverter reading in the log book", style="List Bullet")
d.add_heading("1.1 Cleaning steps", level=2)
d.add_paragraph("Switch off the inverter", style="List Number")
d.add_paragraph("Rinse the panels from the top down", style="List Number")
d.add_paragraph("Switch the inverter on and check the green light", style="List Number")
d.add_heading("2. Who does what", level=1)
d.add_paragraph("Table 1: Responsibilities", style="Caption")
t = d.add_table(rows=4, cols=3)
t.style = "Table Grid"
for i, row in enumerate([["Task", "Person", "How often"], ["Cleaning", "Caretaker", "Monthly"],
                         ["Inverter check", "District engineer", "Quarterly"], ["Log review", "Head teacher", "Monthly"]]):
    for j, v in enumerate(row):
        t.rows[i].cells[j].text = v
d.add_paragraph("The next table shows the fault codes; the first column is shared by two rows.")
m = d.add_table(rows=3, cols=3)
m.style = "Table Grid"
vals = [["Code", "Meaning", "Action"], ["E1", "Low voltage", "Wait for sunlight"], ["", "No voltage", "Call the engineer"]]
for i, row in enumerate(vals):
    for j, v in enumerate(row):
        m.rows[i].cells[j].text = v
m.rows[1].cells[0].merge(m.rows[2].cells[0])
d.add_paragraph("Table 2: Fault codes", style="Caption")
d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
d.add_heading("3. A clean installation", level=1)
d.add_picture(image, width=Cm(10))
d.add_paragraph("Figure 1: Panels after cleaning", style="Caption")
d.add_paragraph("Panels should look like this after every cleaning. ", style="Normal").add_run("Dust cuts output by a fifth.").italic = True
d.add_paragraph("A clean panel is a working panel.", style="Quote")
d.save(out)
print("wrote", out)
